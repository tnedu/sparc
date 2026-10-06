from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import TeamMember
from app.services.access_control import AuthenticatedUser, can_view_product_office, role_capabilities
from app.services.costs import calculate_cost, round_hours
from app.services.estimation_policy import reported_value_rows
from app.services.slugs import slugify

REPORT_DIMENSIONS = {
    "person": "Person",
    "role": "Role",
    "employment_type": "Employment Type",
    "team": "Team",
    "product": "Product",
    "bucket": "Bucket",
}

REPORT_NUMERIC_SORTS = {
    "forecast_cost",
    "actual_cost",
    "variance_cost",
    "forecast_hours",
    "actual_hours",
}


def normalize_labor_cost_dimensions(
    lead: str,
    second: str | None = None,
    third: str | None = None,
    fourth: str | None = None,
) -> list[str]:
    dimensions: list[str] = []
    for raw_dimension in (lead, second, third, fourth):
        dimension = (raw_dimension or "").strip().lower()
        if not dimension or dimension == "none":
            continue
        if dimension not in REPORT_DIMENSIONS:
            raise ValueError(f"Unknown report dimension: {raw_dimension}")
        if dimension not in dimensions:
            dimensions.append(dimension)
    if not dimensions:
        raise ValueError("At least one report dimension is required")
    return dimensions


def normalize_labor_cost_metric(metric: str | None) -> str:
    normalized = (metric or "forecast_cost").strip().lower()
    if normalized not in REPORT_NUMERIC_SORTS and normalized not in REPORT_DIMENSIONS:
        raise ValueError(f"Unknown report metric: {metric}")
    return normalized


def build_labor_cost_report(
    db: Session,
    fiscal_year: int,
    *,
    dimensions: list[str],
    sort_metric: str = "forecast_cost",
    user: AuthenticatedUser | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
) -> dict[str, object]:
    start_month, end_month = _period_month_bounds(fiscal_year, start_date, end_date)
    normalized_dimensions = normalize_labor_cost_dimensions(*dimensions)
    capabilities = role_capabilities(user.role) if user is not None else {}
    if user is not None and not capabilities.get("can_view_reports", False):
        raise PermissionError("Reports are not available for this role")
    if user is not None and "person" in normalized_dimensions and not capabilities.get("can_view_named_people", False):
        raise PermissionError("Person-level reports are not available for this role")
    can_link_people = user is None or capabilities.get("can_view_team_member_profiles", False)
    can_link_teams = user is None or capabilities.get("can_view_team_pages", False)
    normalized_sort_metric = normalize_labor_cost_metric(sort_metric)
    source_rows = [
        row
        for row in reported_value_rows(db, fiscal_year)
        if start_month <= int(row["month_sequence"]) <= end_month
        and (user is None or can_view_product_office(user, str(row["program_area"]) if row["program_area"] else None))
    ]
    member_ids = {int(row["team_member_id"]) for row in source_rows}
    members = {member.id: member for member in db.scalars(select(TeamMember).where(TeamMember.id.in_(member_ids))).all()} if member_ids else {}

    grouped: dict[tuple[tuple[str, str, str | None], ...], dict[str, object]] = {}
    for source_row in source_rows:
        member = members.get(int(source_row["team_member_id"]))
        dimension_values = tuple(
            _dimension_value(
                dimension,
                source_row,
                member,
                can_link_people=can_link_people,
                can_link_teams=can_link_teams,
            )
            for dimension in normalized_dimensions
        )
        if dimension_values not in grouped:
            grouped[dimension_values] = {
                "dimension_values": [
                    {"key": dimension, "label": label, "href": href}
                    for dimension, label, href in dimension_values
                ],
                "forecast_hours": 0.0,
                "actual_hours": 0.0,
                "forecast_cost": 0.0,
                "actual_cost": 0.0,
                "variance_cost": 0.0,
            }

        bill_rate = float(member.bill_rate) if member is not None and member.bill_rate is not None else 0.0
        forecast_hours = float(source_row["forecast_hours"] or 0)
        actual_hours = float(source_row["actual_hours"] or 0)
        forecast_cost = calculate_cost(forecast_hours, bill_rate)
        actual_cost = calculate_cost(actual_hours, bill_rate)
        grouped_row = grouped[dimension_values]
        grouped_row["forecast_hours"] = float(grouped_row["forecast_hours"]) + forecast_hours
        grouped_row["actual_hours"] = float(grouped_row["actual_hours"]) + actual_hours
        grouped_row["forecast_cost"] = float(grouped_row["forecast_cost"]) + forecast_cost
        grouped_row["actual_cost"] = float(grouped_row["actual_cost"]) + actual_cost

    rows = []
    totals = defaultdict(float)
    for grouped_row in grouped.values():
        grouped_row["forecast_hours"] = round_hours(grouped_row["forecast_hours"])
        grouped_row["actual_hours"] = round_hours(grouped_row["actual_hours"])
        grouped_row["forecast_cost"] = round(float(grouped_row["forecast_cost"]), 2)
        grouped_row["actual_cost"] = round(float(grouped_row["actual_cost"]), 2)
        grouped_row["variance_cost"] = round(float(grouped_row["actual_cost"]) - float(grouped_row["forecast_cost"]), 2)
        for metric in ("forecast_hours", "actual_hours", "forecast_cost", "actual_cost", "variance_cost"):
            totals[metric] += float(grouped_row[metric])
        rows.append(grouped_row)

    if normalized_sort_metric in REPORT_DIMENSIONS:
        rows.sort(key=lambda row: (_dimension_sort_value(row, normalized_sort_metric), _dimension_sort_labels(row)))
    else:
        rows.sort(
            key=lambda row: (
                abs(float(row[normalized_sort_metric])),
                _dimension_sort_labels(row),
            ),
            reverse=True,
        )

    return {
        "fiscal_year": fiscal_year,
        "period_start": start_date.isoformat() if start_date else None,
        "period_end": end_date.isoformat() if end_date else None,
        "dimensions": [{"key": dimension, "label": REPORT_DIMENSIONS[dimension]} for dimension in normalized_dimensions],
        "sort_metric": normalized_sort_metric,
        "rows": rows,
        "totals": {
            "forecast_hours": round_hours(totals["forecast_hours"]),
            "actual_hours": round_hours(totals["actual_hours"]),
            "forecast_cost": round(totals["forecast_cost"], 2),
            "actual_cost": round(totals["actual_cost"], 2),
            "variance_cost": round(totals["variance_cost"], 2),
        },
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


def build_labor_cost_report_workbook(
    db: Session,
    fiscal_year: int,
    *,
    dimensions: list[str],
    sort_metric: str = "forecast_cost",
    user: AuthenticatedUser | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
) -> BytesIO:
    report = build_labor_cost_report(
        db,
        fiscal_year,
        dimensions=dimensions,
        sort_metric=sort_metric,
        user=user,
        start_date=start_date,
        end_date=end_date,
    )
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Labor Cost"

    dimensions_label = " > ".join(str(dimension["label"]) for dimension in report["dimensions"])
    worksheet.append(["SPARC Labor Cost Report"])
    period_label = f"{report['period_start']} through {report['period_end']} (fiscal months)" if report["period_start"] else f"FY{fiscal_year}"
    worksheet.append([period_label, f"Dimensions: {dimensions_label}", f"Sorted by: {_metric_label(str(report['sort_metric']))}"])
    worksheet.append([])

    headers = [str(dimension["label"]) for dimension in report["dimensions"]]
    headers.extend(["Forecast Hours", "Forecast Cost", "Actual Hours", "Actual Cost", "Variance Cost"])
    worksheet.append(headers)

    for row in report["rows"]:
        values = [value["label"] for value in row["dimension_values"]]
        values.extend([row["forecast_hours"], row["forecast_cost"], row["actual_hours"], row["actual_cost"], row["variance_cost"]])
        worksheet.append(values)

    worksheet.append([])
    total_row = ["Total"]
    total_row.extend([""] * (len(report["dimensions"]) - 1))
    totals = report["totals"]
    total_row.extend([totals["forecast_hours"], totals["forecast_cost"], totals["actual_hours"], totals["actual_cost"], totals["variance_cost"]])
    worksheet.append(total_row)

    _style_labor_cost_sheet(worksheet, len(headers))

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _dimension_value(
    dimension: str,
    row: dict[str, object],
    member: TeamMember | None,
    *,
    can_link_people: bool,
    can_link_teams: bool,
) -> tuple[str, str, str | None]:
    if dimension == "person":
        href = f"/team-members/{row['team_member_slug']}" if can_link_people else None
        return ("person", str(row["team_member"]), href)
    if dimension == "role":
        role = member.role if member is not None and member.role else "Unassigned"
        return ("role", role, None)
    if dimension == "employment_type":
        employment_type = member.employment_type if member is not None and member.employment_type else "Unassigned"
        return ("employment_type", employment_type, None)
    if dimension == "team":
        team = member.team if member is not None and member.team else "Unassigned"
        href = f"/teams/{slugify(team, fallback='unassigned')}" if can_link_teams else None
        return ("team", team, href)
    if dimension == "product":
        return ("product", str(row["product"]), f"/products/{row['product_slug']}")
    if dimension == "bucket":
        return ("bucket", str(row["bucket"]), None)
    raise ValueError(f"Unknown report dimension: {dimension}")


def _metric_label(metric: str) -> str:
    labels = {
        **REPORT_DIMENSIONS,
        "forecast_cost": "Forecast Cost",
        "actual_cost": "Actual Cost",
        "variance_cost": "Variance Cost",
        "forecast_hours": "Forecast Hours",
        "actual_hours": "Actual Hours",
    }
    return labels[metric]


def _dimension_sort_value(row: dict[str, object], dimension: str) -> str:
    values = row["dimension_values"]
    if not isinstance(values, list):
        return ""
    for value in values:
        if isinstance(value, dict) and value.get("key") == dimension:
            return str(value.get("label") or "").casefold()
    return ""


def _dimension_sort_labels(row: dict[str, object]) -> list[str]:
    values = row["dimension_values"]
    if not isinstance(values, list):
        return []
    return [str(value.get("label") or "").casefold() for value in values if isinstance(value, dict)]


def _period_month_bounds(fiscal_year: int, start_date: date | None, end_date: date | None) -> tuple[int, int]:
    if (start_date is None) != (end_date is None):
        raise ValueError("Both start date and end date are required for a custom report period")
    if start_date is None or end_date is None:
        return 1, 12
    fiscal_start = date(fiscal_year - 1, 7, 1)
    fiscal_end = date(fiscal_year, 6, 30)
    if start_date > end_date:
        raise ValueError("Start date must be on or before end date")
    if start_date < fiscal_start or end_date > fiscal_end:
        raise ValueError(f"Dates must fall within FY{fiscal_year} ({fiscal_start.isoformat()} through {fiscal_end.isoformat()})")
    return _fiscal_month_sequence(start_date), _fiscal_month_sequence(end_date)


def _fiscal_month_sequence(value: date) -> int:
    return ((value.month - 7) % 12) + 1


def _style_labor_cost_sheet(worksheet, column_count: int) -> None:
    header_fill = PatternFill("solid", fgColor="E9EDF5")
    title_font = Font(bold=True, size=14, color="002D72")
    header_font = Font(bold=True, color="76777A")
    total_font = Font(bold=True, color="002D72")

    worksheet["A1"].font = title_font
    worksheet["A2"].font = Font(color="76777A")
    worksheet.freeze_panes = "A5"

    for cell in worksheet[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left")

    first_metric_column = column_count - 4
    for row in worksheet.iter_rows(min_row=5, max_row=worksheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top")
        for cell in row[first_metric_column - 1 :]:
            cell.alignment = Alignment(horizontal="right", vertical="top")

    for row in worksheet.iter_rows(min_row=5, max_row=worksheet.max_row):
        for cell in row:
            if row[0].value == "Total":
                cell.font = total_font
        for cell in (row[first_metric_column - 1], row[first_metric_column + 1]):
            cell.number_format = '#,##0.0'
        for cell in (row[first_metric_column], row[first_metric_column + 2], row[first_metric_column + 3]):
            cell.number_format = '$#,##0;[Red]-$#,##0'

    for index in range(1, column_count + 1):
        column_letter = worksheet.cell(row=4, column=index).column_letter
        worksheet.column_dimensions[column_letter].width = 18 if index <= column_count - 5 else 16
